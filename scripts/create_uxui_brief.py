# -*- coding: utf-8 -*-
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "Livrables"
OUTPUT = OUTPUT_DIR / "BRIEF_CADRAGE_UX_UI_MULTIPRODUIT_MALI.docx"
LOGO = ROOT / "Questionnaire séance 1" / "aee68183-47dd-4bf3-8e9f-f817a23bf7b5.png"

# standard_business_brief with a named MPM brand-color override.
GREEN = "165B3A"
GREEN_DARK = "0F3E29"
GREEN_LIGHT = "EAF4EE"
ORANGE = "F28A2B"
ORANGE_LIGHT = "FFF2E5"
INK = "1C2722"
MUTED = "5E6B65"
LIGHT_GRAY = "F3F5F4"
MID_GRAY = "D9DEDB"
WHITE = "FFFFFF"
YELLOW = "FFF2A8"
RED = "A63B32"

PAGE_WIDTH_DXA = 12240
PAGE_HEIGHT_DXA = 15840
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120


def rgb(hex_color):
    return RGBColor.from_string(hex_color)


def set_font(run, name="Calibri", size=11, color=INK, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:cs"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        tag = "w:" + key
        node = tc_mar.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, color=MID_GRAY, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "start", "bottom", "end", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_table_geometry(table, widths_dxa, indent_dxa=TABLE_INDENT_DXA):
    assert sum(widths_dxa) == CONTENT_WIDTH_DXA, widths_dxa
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[min(idx, len(widths_dxa) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(width / 1440)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_keep_with_next(paragraph, value=True):
    paragraph.paragraph_format.keep_with_next = value


def set_keep_together(paragraph, value=True):
    paragraph.paragraph_format.keep_together = value


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char_1, instr_text, fld_char_2])
    set_font(run, size=8.5, color=MUTED)


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rel_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), GREEN)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.extend([color, underline])
    new_run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    new_run.append(text_node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def add_body(doc, text="", bold_prefix=None, highlight=False, italic=False, after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.10
    p.paragraph_format.widow_control = True
    if bold_prefix and text.startswith(bold_prefix):
        first = p.add_run(bold_prefix)
        set_font(first, bold=True)
        rest = p.add_run(text[len(bold_prefix):])
        set_font(rest, italic=italic)
        if highlight:
            for run in (first, rest):
                run.font.highlight_color = 7
    else:
        run = p.add_run(text)
        set_font(run, italic=italic)
        if highlight:
            run.font.highlight_color = 7
    return p


def add_bullet(doc, text, level=0, bold_prefix=None, highlight=False):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.left_indent = Inches(0.5 + level * 0.25)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.167
    p.paragraph_format.widow_control = True
    if bold_prefix and text.startswith(bold_prefix):
        first = p.add_run(bold_prefix)
        set_font(first, bold=True)
        second = p.add_run(text[len(bold_prefix):])
        set_font(second)
        runs = (first, second)
    else:
        run = p.add_run(text)
        set_font(run)
        runs = (run,)
    if highlight:
        for run in runs:
            run.font.highlight_color = 7
    return p


def paragraph_num_id(paragraph):
    p_pr = paragraph._p.pPr
    if p_pr is None:
        return None
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        return None
    num_id = num_pr.find(qn("w:numId"))
    if num_id is None:
        return None
    return int(num_id.get(qn("w:val")))


def create_decimal_numbering(doc, start=1):
    numbering = doc.part.numbering_part.element
    abstract_ids = [
        int(node.get(qn("w:abstractNumId")))
        for node in numbering.findall(qn("w:abstractNum"))
    ]
    num_ids = [
        int(node.get(qn("w:numId")))
        for node in numbering.findall(qn("w:num"))
    ]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)

    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start_node = OxmlElement("w:start")
    start_node.set(qn("w:val"), str(start))
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "decimal")
    level_text = OxmlElement("w:lvlText")
    level_text.set(qn("w:val"), "%1.")
    level_jc = OxmlElement("w:lvlJc")
    level_jc.set(qn("w:val"), "left")
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "720")
    ind.set(qn("w:hanging"), "360")
    p_pr.extend([tabs, ind])
    level.extend([start_node, num_fmt, level_text, level_jc, p_pr])
    abstract.append(level)
    first_num = numbering.find(qn("w:num"))
    if first_num is None:
        numbering.append(abstract)
    else:
        numbering.insert(list(numbering).index(first_num), abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def add_number(doc, text, level=0, bold_prefix=None):
    previous = doc.paragraphs[-1] if doc.paragraphs else None
    num_id = paragraph_num_id(previous) if previous is not None else None
    if num_id is None:
        num_id = create_decimal_numbering(doc)

    p = doc.add_paragraph()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_node = OxmlElement("w:numId")
    num_id_node.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_id_node])
    p._p.get_or_add_pPr().append(num_pr)
    p.paragraph_format.left_indent = Inches(0.5 + level * 0.25)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.167
    if bold_prefix and text.startswith(bold_prefix):
        first = p.add_run(bold_prefix)
        set_font(first, bold=True)
        second = p.add_run(text[len(bold_prefix):])
        set_font(second)
    else:
        run = p.add_run(text)
        set_font(run)
    return p


def add_h1(doc, text, page_break=False):
    p = doc.add_paragraph(style="Heading 1")
    if page_break:
        p.paragraph_format.page_break_before = True
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    set_font(run, size=16, color=GREEN, bold=True)
    return p


def add_h2(doc, text):
    p = doc.add_paragraph(style="Heading 2")
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    set_font(run, size=13, color=GREEN, bold=True)
    return p


def add_h3(doc, text):
    p = doc.add_paragraph(style="Heading 3")
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    set_font(run, size=11.5, color=GREEN_DARK, bold=True)
    return p


def add_kicker(doc, text, color=ORANGE, after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    run = p.add_run(text.upper())
    set_font(run, size=9, color=color, bold=True)
    return p


def add_callout(doc, label, text, fill=GREEN_LIGHT, accent=GREEN, highlight=False):
    table = doc.add_table(rows=1, cols=1)
    set_repeat_table_header(table.rows[0])
    set_table_geometry(table, [CONTENT_WIDTH_DXA])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_border(cell, color=accent, size="8")
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.08
    r1 = p.add_run(label.upper() + "\n")
    set_font(r1, size=9, color=accent, bold=True)
    r2 = p.add_run(text)
    set_font(r2, size=10.5, color=INK, bold=False)
    if highlight:
        r2.font.highlight_color = 7
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    return table


def add_info_table(doc, headers, rows, widths_dxa, header_fill=GREEN):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths_dxa)
    set_repeat_table_header(table.rows[0])
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, header_fill)
        set_cell_border(cell, color=WHITE, size="5")
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(header)
        set_font(run, size=9.5, color=WHITE, bold=True)
    for row_idx, values in enumerate(rows):
        cells = table.add_row().cells
        for idx, value in enumerate(values):
            cell = cells[idx]
            set_cell_shading(cell, WHITE if row_idx % 2 == 0 else LIGHT_GRAY)
            set_cell_border(cell)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            run = p.add_run(str(value))
            set_font(run, size=9.2, color=INK)
    set_table_geometry(table, widths_dxa)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(2)
    return table


def add_code_block(doc, code, title=None):
    table = doc.add_table(rows=2 if title else 1, cols=1)
    set_repeat_table_header(table.rows[0])
    set_table_geometry(table, [CONTENT_WIDTH_DXA])
    if title:
        header_cell = table.cell(0, 0)
        set_cell_shading(header_cell, GREEN_LIGHT)
        set_cell_border(header_cell, color=MID_GRAY, size="5")
        set_cell_margins(header_cell, top=70, start=160, bottom=70, end=160)
        header_p = header_cell.paragraphs[0]
        header_p.paragraph_format.space_before = Pt(0)
        header_p.paragraph_format.space_after = Pt(0)
        header_run = header_p.add_run(title)
        set_font(header_run, size=9.5, color=GREEN_DARK, bold=True)
        cell = table.cell(1, 0)
    else:
        cell = table.cell(0, 0)
    set_cell_shading(cell, "F6F7F6")
    set_cell_border(cell, color=MID_GRAY, size="5")
    set_cell_margins(cell, top=120, start=160, bottom=120, end=160)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    lines = code.strip("\n").split("\n")
    for index, line in enumerate(lines):
        run = p.add_run(line)
        set_font(run, name="Consolas", size=7.6, color="24312B")
        if index < len(lines) - 1:
            run.add_break()
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(2)


def add_logo_reference(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    shape = p.add_run().add_picture(str(LOGO), width=Inches(2.25), height=Inches(2.0))
    shape._inline.docPr.set("descr", "Logo raster actuel de Multiproduit Mali SARL, à vectoriser")
    blip_fill = shape._inline.graphic.graphicData.pic.blipFill
    src_rect = OxmlElement("a:srcRect")
    src_rect.set("l", "5000")
    src_rect.set("t", "28000")
    src_rect.set("r", "10000")
    src_rect.set("b", "30000")
    blip_fill.insert(1, src_rect)
    p.paragraph_format.space_after = Pt(3)
    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(8)
    run = caption.add_run("Logo source actuellement disponible - référence de reconstruction, non livrable final")
    set_font(run, size=8.5, color=MUTED, italic=True)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.78)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.36)
    section.footer_distance = Inches(0.36)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    normal.paragraph_format.widow_control = True

    for style_name, size, color, before, after in (
        ("Heading 1", 16, GREEN, 16, 8),
        ("Heading 2", 13, GREEN, 12, 6),
        ("Heading 3", 11.5, GREEN_DARK, 8, 4),
    ):
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.widow_control = True

    for style_name in ("List Bullet", "List Bullet 2", "List Number", "List Number 2"):
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(11)
        style.font.color.rgb = rgb(INK)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.167

    header = section.header
    hp = header.paragraphs[0]
    hp.paragraph_format.space_after = Pt(0)
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hr = hp.add_run("MULTIPRODUIT MALI  |  BRIEF UX/UI")
    set_font(hr, size=8, color=MUTED, bold=True)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.paragraph_format.space_after = Pt(0)
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fr = fp.add_run("Document de consultation - EKAM Capital  |  Page ")
    set_font(fr, size=8.5, color=MUTED)
    add_page_number(fp)

    core = doc.core_properties
    core.title = "Brief de cadrage UX/UI - Multiproduit Mali"
    core.subject = "Vitrine B2B internationale, produits, partenaires et communauté bilingue"
    core.author = "EKAM Capital"
    core.keywords = "Multiproduit Mali, UX, UI, motion design, partenaires, boissons, bilingue"
    core.comments = "Document de consultation préparé pour l'orientation UX/UI du site multiproduitmali.ml"


def build_document():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_document(doc)

    # Cover - customer_pack pattern.
    add_kicker(doc, "Brief de cadrage UX/UI et motion design", color=ORANGE, after=6)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("MULTIPRODUIT\nMALI")
    set_font(r, size=31, color=GREEN_DARK, bold=True)
    set_keep_together(p)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(20)
    r = p.add_run("Vitrine B2B internationale et expérience digitale de marques")
    set_font(r, size=15, color=MUTED, bold=False)

    table = doc.add_table(rows=2, cols=2)
    set_repeat_table_header(table.rows[0])
    set_table_geometry(table, [4680, 4680])
    meta = [
        ("Préparé pour", "Professionnel UX/UI et motion designer"),
        ("Préparé par", "EKAM Capital - Mahamadou Karamoko DIARRA, Représentant Mali"),
        ("Domaine prévu", "multiproduitmali.ml"),
        ("Version", "1.0 - 25 août 2026"),
    ]
    for idx, (label, value) in enumerate(meta):
        row = idx // 2
        col = idx % 2
        cell = table.cell(row, col)
        set_cell_shading(cell, GREEN_LIGHT if row == 0 else ORANGE_LIGHT)
        set_cell_border(cell, color=WHITE, size="8")
        cell.paragraphs[0].clear()
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r1 = p.add_run(label.upper() + "\n")
        set_font(r1, size=8.5, color=GREEN, bold=True)
        r2 = p.add_run(value)
        set_font(r2, size=10, color=INK, bold=True)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)
    add_callout(
        doc,
        "Décision de cadrage",
        "Multiproduit Mali SARL est l'identité institutionnelle unique. Le nom provisoire SIMPARA ne doit apparaître nulle part dans l'expérience finale. Tropicoul et Triplex sont présentés comme marques ou gammes de produits.",
        fill=GREEN_LIGHT,
        accent=GREEN,
    )

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("Objet : obtenir une orientation UX/UI précise, argumentée et directement exploitable par l'équipe de développement.")
    set_font(r, size=10.5, color=INK, italic=True)
    p.add_run().add_break(WD_BREAK.PAGE)

    # Executive summary.
    add_h1(doc, "Synthèse décisionnelle")
    add_body(
        doc,
        "Le futur site n'est pas seulement un catalogue de boissons. Il doit établir Multiproduit Mali comme une entreprise crédible, moderne et prête à développer des partenariats de distribution dans plusieurs pays, tout en conservant une expression visuelle énergique et populaire autour de ses produits.",
    )
    add_callout(
        doc,
        "Direction recommandée",
        "Une expérience composée à 60 % de crédibilité institutionnelle et à 40 % de spectacle produit. Le hero adopte une parade fluide de produits; une carte vivante soutient le recrutement international; chaque page produit développe son propre univers tropical.",
        fill=ORANGE_LIGHT,
        accent=ORANGE,
    )
    add_info_table(
        doc,
        ["Élément", "Décision de cadrage"],
        [
            ("Marque principale", "Multiproduit Mali SARL / MPM"),
            ("Cibles prioritaires", "Partenaires, importateurs, distributeurs et représentants par pays"),
            ("Cible secondaire", "Consommateurs, visiteurs des actualités et communauté"),
            ("Langues", "Français et anglais, sélection automatique avec choix manuel persistant"),
            ("Produit d'appel", "Boissons fruitées Tropicoul et boisson énergisante Triplex"),
            ("Interaction phare", "Produits en mouvement fluide dans le hero; 3D ou rendu 2.5D optimisé"),
            ("Conversion principale", "Dépôt d'une candidature pour représenter ou distribuer les produits"),
        ],
        [2600, 6760],
    )
    add_h2(doc, "Le mandat attendu du professionnel UX/UI")
    add_bullet(doc, "Interroger et renforcer le positionnement, sans revenir au nom provisoire SIMPARA.")
    add_bullet(doc, "Définir une architecture claire pour un site B2B international, sans sacrifier l'attractivité des produits.")
    add_bullet(doc, "Proposer des wireframes desktop et mobile, un système visuel et un langage de mouvement mesurables.")
    add_bullet(doc, "Prévoir les états réels : chargement 3D, erreurs de formulaire, traduction, modération, faible débit et réduction des animations.")
    add_bullet(doc, "Produire une documentation de handoff suffisamment précise pour éviter une interprétation libre par le développeur.")

    add_h1(doc, "Sommaire et mode d'emploi", page_break=True)
    for item in (
        "1. Contexte, marque et objectifs",
        "2. Publics prioritaires et parcours",
        "3. Architecture de l'information",
        "4. Direction artistique, hero et motion design",
        "5. Produits, 3D et pages immersives",
        "6. Espace partenaires et formulaire",
        "7. Bilingue, géolocalisation et communauté",
        "8. Responsive, accessibilité et performance",
        "9. Logo, réseaux sociaux et interface de gestion",
        "10. Livrables, critères de réussite et points à confirmer",
        "Annexe A. Prompt complet pour le professionnel UX/UI",
        "Annexe B. Orientations techniques et extraits de code",
        "Annexe C. Références",
    ):
        add_bullet(doc, item)
    add_callout(
        doc,
        "Convention de lecture",
        "Les éléments surlignés en jaune sont des informations manquantes ou des arbitrages à obtenir avant la conception haute fidélité ou le développement final.",
        fill=YELLOW,
        accent=ORANGE,
    )

    # 1. Context.
    add_h1(doc, "1. Contexte, marque et objectifs", page_break=True)
    add_h2(doc, "1.1 Contexte du projet")
    add_body(
        doc,
        "Multiproduit Mali SARL souhaite disposer d'un site vitrine dynamique destiné à présenter son entreprise, valoriser ses produits et ouvrir des opportunités de représentation ou de distribution dans d'autres pays. Le domaine prévu est multiproduitmali.ml.",
    )
    add_body(
        doc,
        "Le projet était initialement désigné sous le nom SIMPARA Distribution. Ce nom était provisoire et doit être retiré de tous les écrans, URL internes, métadonnées, courriels, textes alternatifs et contenus de démonstration. L'entreprise à présenter est Multiproduit Mali SARL.",
    )
    add_h2(doc, "1.2 Offre et contenu connus")
    add_bullet(doc, "Gamme Tropicoul : boissons fruitées associées notamment à l'ananas, l'orange, la mangue, la goyave, le cocktail et le tamarin.")
    add_bullet(doc, "Triplex : boisson énergisante avec un univers plus sportif, électrique et urbain.")
    add_bullet(doc, "Actualités : activités de l'entreprise, événements, opérations commerciales et contenus de terrain.")
    add_bullet(doc, "Communauté : réactions du public, commentaires et traduction français-anglais.")
    add_bullet(doc, "Partenariats : candidature de distributeurs, représentants et partenaires commerciaux.")
    add_bullet(doc, "Interface de gestion : produits, contenus, commentaires, traductions et candidatures.")
    add_body(doc, "À CONFIRMER : liste commerciale définitive, formats de conditionnement, volumes, composition, arguments de vente et disponibilité par pays.", highlight=True)

    add_h2(doc, "1.3 Objectifs métier")
    add_number(doc, "Inspirer confiance à un partenaire potentiel dans les cinq premières secondes.")
    add_number(doc, "Démontrer la qualité, la variété et le potentiel commercial du portefeuille.")
    add_number(doc, "Transformer l'intérêt en candidature partenaire structurée.")
    add_number(doc, "Faciliter la découverte des activités de l'entreprise et renforcer sa présence sociale.")
    add_number(doc, "Servir un public international en français et en anglais.")
    add_number(doc, "Permettre à l'équipe MPM d'administrer le site sans intervention technique courante.")

    add_h2(doc, "1.4 Indicateurs à prévoir")
    add_info_table(
        doc,
        ["Objectif", "Indicateur UX ou métier"],
        [
            ("Crédibilité", "Compréhension de l'activité et de la provenance lors de tests utilisateurs courts"),
            ("Partenariat", "Taux de clic vers Devenir partenaire et taux d'achèvement du formulaire"),
            ("Produits", "Consultation des gammes et interaction avec le sélecteur produit"),
            ("International", "Usage du sélecteur de langue et origine des candidatures"),
            ("Communauté", "Commentaires approuvés, traductions demandées et réactions"),
            ("Performance", "LCP, INP, CLS, poids 3D et fluidité sur mobile réel"),
        ],
        [3000, 6360],
    )

    # 2. Audiences.
    add_h1(doc, "2. Publics prioritaires et parcours")
    add_h2(doc, "2.1 Partenaire de distribution international - priorité 1")
    add_body(doc, "Profil : importateur, grossiste, chaîne de distribution ou entreprise locale cherchant une gamme de boissons à développer sur son marché.")
    add_bullet(doc, "Questions : l'entreprise est-elle sérieuse? Quels produits propose-t-elle? Quelle capacité de production et quel accompagnement commercial offre-t-elle?")
    add_bullet(doc, "Preuves attendues : histoire, qualité, portefeuille, marchés, coordonnées professionnelles, actualités et processus de partenariat.")
    add_bullet(doc, "Action principale : consulter le portefeuille puis soumettre une candidature qualifiée.")

    add_h2(doc, "2.2 Représentant potentiel dans un pays - priorité 2")
    add_body(doc, "Profil : entrepreneur, agent commercial ou structure disposant d'un réseau local et souhaitant représenter MPM.")
    add_bullet(doc, "Besoin : comprendre rapidement les conditions d'intérêt, le territoire recherché et les informations à fournir.")
    add_bullet(doc, "Action principale : compléter le formulaire, joindre une présentation et recevoir une confirmation.")

    add_h2(doc, "2.3 Consommateur et membre de la communauté - priorité 3")
    add_body(doc, "Profil : personne découvrant un produit, suivant une activité ou souhaitant réagir à une publication.")
    add_bullet(doc, "Besoin : expérience mobile rapide, visuelle, sociale et compréhensible dans sa langue.")
    add_bullet(doc, "Action principale : découvrir un produit, suivre un réseau social, commenter ou traduire un commentaire.")

    add_h2(doc, "2.4 Parcours de référence")
    add_info_table(
        doc,
        ["Parcours", "Étapes minimales"],
        [
            ("Partenaire", "Accueil > crédibilité MPM > portefeuille > avantages > candidature > confirmation"),
            ("Représentant", "Accueil ou lien social > Devenir partenaire > territoire > profil > pièce jointe > confirmation"),
            ("Produit", "Accueil > produit actif > page produit > détails > autre parfum ou contact"),
            ("Communauté", "Actualité > lecture > commentaires > traduction > réaction ou publication"),
        ],
        [2500, 6860],
    )

    # 3. IA.
    add_h1(doc, "3. Architecture de l'information")
    add_h2(doc, "3.1 Navigation principale")
    for item in (
        "Accueil",
        "Découvrir Multiproduit Mali",
        "Nos produits",
        "Devenir partenaire",
        "Actualités et communauté",
        "Contact",
        "Sélecteur FR / EN",
    ):
        add_bullet(doc, item)
    add_body(doc, "La navigation doit faire de Devenir partenaire une commande visible, sans transformer tous les liens en boutons concurrents.")

    add_h2(doc, "3.2 Arborescence recommandée")
    add_info_table(
        doc,
        ["Page", "Contenu et fonction"],
        [
            ("Accueil", "Promesse, produits, crédibilité, partenariat, marchés, actualités, réseaux sociaux"),
            ("L'entreprise", "Histoire, mission, savoir-faire, valeurs, présence, preuves et équipe si validée"),
            ("Produits", "Vue d'ensemble par gamme, filtres simples et accès aux pages détaillées"),
            ("Page produit", "Univers visuel, contenant, parfum, bénéfices factuels, formats et contact"),
            ("Partenaires", "Proposition B2B, profils recherchés, processus et formulaire"),
            ("Actualités", "Publications administrables, médias, réactions et commentaires"),
            ("Contact", "Coordonnées, réseaux, carte ou adresse si publiée, formulaire général"),
            ("Mentions", "Confidentialité, cookies, conditions d'utilisation de la communauté"),
            ("Administration", "Accès privé aux contenus, produits, candidatures, commentaires et traductions"),
        ],
        [2200, 7160],
    )

    add_h2(doc, "3.3 Composition de la page d'accueil")
    steps = [
        ("Hero", "Parade dynamique des produits, promesse internationale et deux actions."),
        ("Preuve immédiate", "Bandeau court : identité MPM, activité, origine et présence."),
        ("Entreprise", "Présentation institutionnelle avec un contenu réel de terrain."),
        ("Portefeuille", "Explorateur interactif Tropicoul / Triplex."),
        ("Partenariat", "Proposition claire : représenter ou distribuer dans son pays."),
        ("Marchés", "Carte vivante des zones actuelles et des opportunités, sans inventer de présence."),
        ("Actualités", "Trois contenus récents avec accès à la communauté."),
        ("Social", "Liens visibles, contenus récents si l'intégration est autorisée."),
        ("Clôture", "Contact, candidature et informations légales."),
    ]
    add_info_table(doc, ["Section", "Rôle"], steps, [2100, 7260])

    # 4. Creative direction.
    add_h1(doc, "4. Direction artistique, hero et motion design")
    add_h2(doc, "4.1 Positionnement visuel")
    add_callout(
        doc,
        "Principe 60 / 40",
        "60 % de crédibilité institutionnelle : structure, preuves, lisibilité et sobriété. 40 % de spectacle produit : couleurs, 3D, ingrédients et mouvement.",
        fill=GREEN_LIGHT,
        accent=GREEN,
    )
    add_bullet(doc, "Adjectifs : malien, ambitieux, généreux, moderne et international.")
    add_bullet(doc, "Palette : vert MPM, orange du logo, blanc, encre sombre et couleurs spécifiques aux produits.")
    add_bullet(doc, "Photographie : produits lisibles, scènes réelles, distribution, équipe et activités; éviter les images génériques de bureau.")
    add_bullet(doc, "Typographie : caractère affirmé pour les titres et excellente lisibilité pour les informations commerciales.")
    add_bullet(doc, "Composition : sections en bandes pleines, grilles aérées et produits en premier plan; éviter les empilements de cartes décoratives.")

    add_h2(doc, "4.2 Trois pistes pour le hero")
    add_info_table(
        doc,
        ["Piste", "Description", "Usage recommandé"],
        [
            ("A. Parade des produits", "Les produits traversent l'écran; celui au centre devient actif et tourne.", "Hero principal - recommandée"),
            ("B. Carte vivante", "Bamako devient le point de départ de routes vers d'autres marchés.", "Section partenaires et marchés"),
            ("C. Explosion tropicale", "Fruits, feuilles, glace et gouttelettes composent l'univers du produit.", "Pages produits"),
        ],
        [2100, 4300, 2960],
    )

    add_h2(doc, "4.3 Storyboard précis du hero recommandé")
    add_number(doc, "À l'ouverture, le logo MPM et la promesse sont visibles sans attendre le chargement 3D.")
    add_number(doc, "Une première canette ou bouteille apparaît au centre, accompagnée des produits voisins partiellement visibles.")
    add_number(doc, "Le produit actif tourne lentement; le fond et les ingrédients changent selon sa gamme.")
    add_number(doc, "Le visiteur contrôle le passage avec le scroll, le glissement tactile, les flèches et le clavier.")
    add_number(doc, "Le bouton Devenir partenaire reste stable; Explorer nos produits est secondaire.")
    add_number(doc, "La fin du hero laisse entrevoir la section Découvrir Multiproduit Mali dans le premier viewport.")

    add_h2(doc, "4.4 Règles chiffrées de mouvement")
    add_info_table(
        doc,
        ["Comportement", "Règle proposée"],
        [
            ("Hauteur du hero", "Environ 86-92 svh, avec un aperçu visible de la section suivante"),
            ("Transition de produit", "650 à 900 ms, accélération puis décélération maîtrisées"),
            ("Rotation au repos", "Une rotation complète lente en 14 à 18 secondes, désactivable"),
            ("Micro-interaction", "160 à 220 ms"),
            ("Transition de section", "300 à 450 ms"),
            ("Parallaxe", "Amplitude maximale de 12 px et environ 2 degrés"),
            ("Couches animées", "Deux couches principales simultanées, hors produit"),
            ("Scroll", "Scroll natif conservé; aucun blocage artificiel de la navigation"),
            ("Réduction du mouvement", "Image statique, carrousel manuel et transitions par fondu"),
        ],
        [3100, 6260],
    )

    add_h2(doc, "4.5 Ce qu'il faut éviter")
    add_bullet(doc, "Un défilement automatique rapide qui empêche la lecture.")
    add_bullet(doc, "Une accumulation de fruits, particules, texte et 3D qui masque le produit.")
    add_bullet(doc, "Des animations déclenchées sur chaque paragraphe.")
    add_bullet(doc, "Des sections épinglées sur plusieurs écrans, particulièrement pénibles sur mobile.")
    add_bullet(doc, "Un hero spectaculaire qui ne dit pas ce que fait Multiproduit Mali.")

    # 5. Products and 3D.
    add_h1(doc, "5. Produits, 3D et pages immersives")
    add_h2(doc, "5.1 Principe de conception")
    add_body(doc, "La structure des pages produits doit rester cohérente, tandis que l'ambiance se transforme progressivement en fonction de la gamme et du parfum. La différence doit venir de la couleur, des ingrédients, de la lumière et du rythme, pas d'une navigation différente à chaque page.")
    add_info_table(
        doc,
        ["Univers", "Expression"],
        [
            ("Tropicoul Ananas", "Jaune franc, feuillage vert, tranches d'ananas, lumière solaire"),
            ("Tropicoul Orange", "Orange lumineux, agrumes, fraîcheur et mouvement circulaire"),
            ("Tropicoul Mangue", "Jaune-orangé, pulpe, chaleur maîtrisée et douceur"),
            ("Tropicoul Goyave", "Rose et vert, feuilles fines, ambiance généreuse"),
            ("Tropicoul Cocktail", "Palette multicolore contrôlée, composition plus festive"),
            ("Tropicoul Tamarin", "Vert profond et accents chauds, univers plus mature"),
            ("Triplex", "Contraste fort, énergie, lumière nette et rythme plus rapide"),
        ],
        [2700, 6660],
    )

    add_h2(doc, "5.2 Stratégie 3D retenue pour la maquette")
    add_body(doc, "Les contenants étant standards et les fichiers d'étiquettes indisponibles, un prototype peut être obtenu à partir de photographies Multi-View dans Meshy AI. Le résultat brut ne doit cependant être validé qu'après inspection de la lisibilité du logo, des raccords de texture et des reflets.")
    add_number(doc, "Photographier une référence sous quatre vues cohérentes : face, gauche, dos et droite.")
    add_number(doc, "Utiliser un éclairage diffus, un appareil fixe et une rotation régulière du produit.")
    add_number(doc, "Générer un premier GLB et contrôler la texture sur 360 degrés.")
    add_number(doc, "Nettoyer la texture et, si nécessaire, transférer celle-ci sur une géométrie cylindrique propre.")
    add_number(doc, "Tester sur un téléphone réel avant de traiter toutes les références.")
    add_callout(
        doc,
        "Critère de validation 3D",
        "À distance normale de consultation, le logo doit être lisible, aucun raccord ne doit attirer l'œil, les reflets ne doivent pas tourner comme s'ils étaient imprimés et le modèle ne doit pas provoquer de ralentissement visible.",
        fill=ORANGE_LIGHT,
        accent=ORANGE,
    )

    add_h2(doc, "5.3 Budget de performance à transmettre au designer 3D")
    add_bullet(doc, "Utiliser GLB/glTF avec compression de géométrie et de textures lorsque le pipeline le permet.")
    add_bullet(doc, "Objectif initial : modèle principal proche de 1 à 3 Mo après optimisation.")
    add_bullet(doc, "Prévoir une image WebP/AVIF de remplacement visible immédiatement.")
    add_bullet(doc, "Limiter le ratio de pixels du canvas sur mobile et arrêter le rendu hors écran.")
    add_bullet(doc, "Partager la géométrie entre variantes lorsqu'un même contenant est utilisé.")

    # 6. Partner form.
    add_h1(doc, "6. Espace partenaires et formulaire")
    add_h2(doc, "6.1 Proposition de valeur")
    add_body(doc, "La page doit répondre à quatre questions : pourquoi collaborer avec MPM, quels profils sont recherchés, comment la relation peut commencer et quelles informations sont nécessaires pour être recontacté.")
    add_callout(
        doc,
        "Message possible",
        "Représentez nos marques dans votre pays. Présentez-nous votre marché, votre réseau et votre projet; notre équipe étudiera votre candidature.",
        fill=GREEN_LIGHT,
        accent=GREEN,
    )

    add_h2(doc, "6.2 Structure du formulaire")
    add_info_table(
        doc,
        ["Groupe", "Champs attendus"],
        [
            ("Identité", "Nom, prénom, fonction, entreprise"),
            ("Coordonnées", "Pays, ville, courriel, téléphone / WhatsApp, langue préférée"),
            ("Projet", "Territoire, produits concernés, type de partenariat"),
            ("Capacité", "Expérience, réseau, canaux de vente, zone couverte"),
            ("Présentation", "Message et pièce jointe facultative"),
            ("Conformité", "Consentement, politique de confidentialité et dispositif anti-spam"),
        ],
        [2200, 7160],
    )
    add_h2(doc, "6.3 Règles UX du formulaire")
    add_bullet(doc, "Deux étapes courtes au maximum sur mobile; une seule page lisible sur ordinateur.")
    add_bullet(doc, "Libellés visibles au-dessus des champs; ne pas utiliser le placeholder comme unique libellé.")
    add_bullet(doc, "Validation au bon moment, avec message d'erreur précis et conservation des données saisies.")
    add_bullet(doc, "Formats de téléphone et de pays internationalisés.")
    add_bullet(doc, "Pièce jointe limitée à des formats et tailles explicitement indiqués.")
    add_bullet(doc, "Écran de succès avec numéro de référence ou confirmation claire.")
    add_bullet(doc, "Courriel interne détaillé et accusé de réception bilingue au candidat.")

    add_h2(doc, "6.4 Suivi dans l'administration")
    add_bullet(doc, "Liste filtrable par pays, date, type de partenariat et statut.")
    add_bullet(doc, "Statuts : nouveau, en étude, contacté, retenu, refusé.")
    add_bullet(doc, "Fiche complète, téléchargement de la pièce jointe et historique minimal.")
    add_bullet(doc, "Export CSV réservé aux administrateurs autorisés.")
    add_body(doc, "À CONFIRMER : adresse de réception des candidatures, taille maximale des pièces jointes, durée de conservation et responsables autorisés.", highlight=True)

    # 7. i18n/community.
    add_h1(doc, "7. Bilingue, géolocalisation et communauté")
    add_h2(doc, "7.1 Règle de sélection de langue")
    add_number(doc, "Choix manuel déjà enregistré dans le navigateur.")
    add_number(doc, "Langue préférée communiquée par le navigateur.")
    add_number(doc, "Pays détecté de manière approximative côté serveur, sans demander la position GPS.")
    add_number(doc, "Français par défaut en l'absence de signal exploitable.")
    add_body(doc, "Le choix manuel FR / EN doit toujours rester visible et persistant. Les drapeaux ne doivent pas remplacer les codes de langue, car une langue ne correspond pas à un seul pays.")

    add_h2(doc, "7.2 Contenu éditorial")
    add_bullet(doc, "Tous les menus, pages, formulaires, courriels et messages système existent en français et en anglais.")
    add_bullet(doc, "Les textes institutionnels et commerciaux sont traduits puis validés humainement.")
    add_bullet(doc, "Chaque contenu administrable possède un statut de traduction : brouillon, à valider, publié.")
    add_bullet(doc, "Les URL sont localisées, par exemple /fr/produits et /en/products, avec liens hreflang.")

    add_h2(doc, "7.3 Commentaires et traduction")
    add_number(doc, "Conserver et afficher le commentaire original avec sa langue détectée.")
    add_number(doc, "Proposer Voir la traduction au lieu de remplacer silencieusement le texte.")
    add_number(doc, "Générer la version française ou anglaise à la demande ou lors de la modération.")
    add_number(doc, "Mettre la traduction en cache afin de ne pas repayer chaque consultation.")
    add_number(doc, "Afficher la mention Traduction automatique et permettre de revenir à l'original.")
    add_number(doc, "Soumettre les commentaires à une modération, un anti-spam et des règles communautaires.")
    add_body(doc, "À CONFIRMER : fournisseur de traduction automatique, budget mensuel, politique de modération et délai de publication.", highlight=True)

    # 8. Responsive/accessibility/performance.
    add_h1(doc, "8. Responsive, accessibilité et performance")
    add_h2(doc, "8.1 Responsive réel")
    add_body(doc, "Le mobile n'est pas la version desktop compressée. Le professionnel doit redéfinir l'ordre, la densité, le mouvement et les contrôles pour chaque largeur utile.")
    add_info_table(
        doc,
        ["Viewport de contrôle", "Vérification attendue"],
        [
            ("360 x 800", "Android compact : lecture, boutons, formulaire et hero"),
            ("390 x 844", "Mobile courant : composition principale et interactions tactiles"),
            ("768 x 1024", "Tablette : navigation, grilles et carte"),
            ("1024 x 768", "Petit ordinateur : hero et sections sans collision"),
            ("1440 x 900", "Desktop : hiérarchie, respiration et scène produit"),
        ],
        [2500, 6860],
    )
    add_bullet(doc, "Cibles tactiles d'au moins 44 x 44 px dans le design interne, même lorsque le minimum réglementaire est inférieur.")
    add_bullet(doc, "Aucune police mise à l'échelle directement avec la largeur du viewport.")
    add_bullet(doc, "Aucun texte ou contrôle ne doit recouvrir le produit, un autre texte ou la navigation.")
    add_bullet(doc, "Les carrousels doivent fonctionner par glissement, boutons et clavier, sans dépendre uniquement du drag.")
    add_bullet(doc, "Respecter les encoches et zones sûres des appareils mobiles.")

    add_h2(doc, "8.2 Accessibilité")
    add_bullet(doc, "Cible de conformité : WCAG 2.2 niveau AA.")
    add_bullet(doc, "Contraste du texte normal d'au moins 4,5:1; états de focus visibles et non masqués.")
    add_bullet(doc, "Navigation complète au clavier; ordre logique du focus; lien d'évitement.")
    add_bullet(doc, "Alternative textuelle aux produits 3D et aux cartes interactives.")
    add_bullet(doc, "Respect de prefers-reduced-motion et possibilité d'arrêter un mouvement automatique prolongé.")
    add_bullet(doc, "Erreurs de formulaire identifiées par texte, pas uniquement par couleur.")

    add_h2(doc, "8.3 Objectifs de performance")
    add_info_table(
        doc,
        ["Mesure", "Cible"],
        [
            ("Largest Contentful Paint", "Inférieur ou égal à 2,5 s au 75e percentile"),
            ("Interaction to Next Paint", "Inférieur ou égal à 200 ms au 75e percentile"),
            ("Cumulative Layout Shift", "Inférieur ou égal à 0,1 au 75e percentile"),
            ("3D", "Chargement différé, visuel de remplacement et arrêt hors écran"),
            ("Images", "Formats modernes, dimensions explicites et chargement différé hors hero"),
        ],
        [3100, 6260],
    )

    # 9. Logo, social, admin.
    add_h1(doc, "9. Logo, réseaux sociaux et interface de gestion")
    add_h2(doc, "9.1 Retouche du logo")
    add_logo_reference(doc)
    add_body(doc, "Le fichier disponible est un raster compressé présentant des contours flous et un élément parasite sur la droite. Un simple agrandissement par IA ne constitue pas une solution de marque. La prestation recommandée est une reconstruction vectorielle fidèle.")
    add_bullet(doc, "Conserver le cercle, les flèches, le sigle MPM, le vert, l'orange et les mentions officielles.")
    add_bullet(doc, "Corriger les espacements, la régularité du texte circulaire, les contours et l'équilibre du symbole.")
    add_bullet(doc, "Prévoir SVG, PDF vectoriel, PNG transparent 2048 px, version monochrome et version inversée.")
    add_bullet(doc, "Créer un monogramme MPM pour favicon et petite taille, ainsi qu'une variante horizontale pour l'en-tête.")
    add_body(doc, "À CONFIRMER : orthographe juridique exacte des mentions du logo, couleurs officielles, police historique et degré de liberté autorisé pour la retouche.", highlight=True)

    add_h2(doc, "9.2 Réseaux sociaux")
    add_bullet(doc, "Icônes visibles dans l'en-tête desktop, le menu mobile, la zone actualités et le pied de page.")
    add_bullet(doc, "WhatsApp peut être une action de contact sur la page partenaires, sans bulle flottante masquant le contenu.")
    add_bullet(doc, "Utiliser les icônes officielles ou une bibliothèque reconnue, avec libellés accessibles.")
    add_bullet(doc, "Les intégrations de flux sociaux ne doivent pas ralentir le hero; utiliser des aperçus maîtrisés si nécessaire.")
    add_body(doc, "À CONFIRMER : URL officielles Facebook, Instagram, TikTok, LinkedIn, YouTube et numéro WhatsApp professionnel.", highlight=True)

    add_h2(doc, "9.3 Interface de gestion")
    add_info_table(
        doc,
        ["Module", "Fonctions minimales"],
        [
            ("Produits", "Créer, modifier, classer, publier, gérer images, 3D et traductions"),
            ("Actualités", "Rédiger, programmer, illustrer et publier en FR / EN"),
            ("Communauté", "Modérer, masquer, traduire et signaler les commentaires"),
            ("Partenaires", "Consulter, filtrer, changer le statut et exporter les candidatures"),
            ("Paramètres", "Coordonnées, réseaux, destinataires d'e-mails et textes légaux"),
            ("Utilisateurs", "Au minimum administrateur et éditeur, selon la solution retenue"),
        ],
        [2300, 7060],
    )

    # 10. Deliverables and technical arbitration.
    add_h1(doc, "10. Livrables, critères de réussite et points à confirmer")
    add_h2(doc, "10.1 Livrables attendus du professionnel UX/UI")
    add_number(doc, "Audit synthétique du cadrage et des risques.")
    add_number(doc, "Arborescence et parcours prioritaires.")
    add_number(doc, "Wireframes basse fidélité desktop et mobile.")
    add_number(doc, "Deux variantes du hero puis recommandation argumentée.")
    add_number(doc, "Maquettes haute fidélité de l'accueil, produits, page produit, partenaires, actualités et administration essentielle.")
    add_number(doc, "Prototype interactif couvrant le hero, la navigation mobile, le formulaire et la traduction d'un commentaire.")
    add_number(doc, "Mini design system : couleurs, typographie, grille, icônes, boutons, formulaires, états et composants.")
    add_number(doc, "Spécification motion : déclencheur, durée, easing, amplitude, comportement mobile et reduced motion.")
    add_number(doc, "Handoff développeur : mesures, tokens, exports, règles responsive et états d'erreur / chargement.")

    add_h2(doc, "10.2 Critères d'acceptation UX/UI")
    add_bullet(doc, "Multiproduit Mali, son activité et la possibilité de partenariat sont compris immédiatement.")
    add_bullet(doc, "Le hero reste lisible avant, pendant et après le chargement du produit 3D.")
    add_bullet(doc, "Le principal appel à l'action n'entre pas en concurrence avec plusieurs boutons équivalents.")
    add_bullet(doc, "Les parcours partenaire et commentaire sont entièrement dessinés, y compris succès, erreurs et modération.")
    add_bullet(doc, "Les versions française et anglaise sont prévues dans les composants, avec les différences de longueur de texte.")
    add_bullet(doc, "Les cinq viewports de contrôle sont présentés et ne contiennent aucun chevauchement.")
    add_bullet(doc, "Les comportements reduced motion, faible débit et absence de WebGL sont documentés.")

    add_h2(doc, "10.3 Arbitrage technique obligatoire avant développement final")
    add_callout(
        doc,
        "Point sensible",
        "La maquette actuelle repose sur React 19, Vinext et des outils Cloudflare. L'offre Hostinger Premium communiquée indique que Node.js n'est pas inclus. Or l'administration, la communauté, la traduction et les candidatures nécessitent une architecture serveur ou des services externes.",
        fill=YELLOW,
        accent=ORANGE,
        highlight=True,
    )
    add_info_table(
        doc,
        ["Option", "Conséquence"],
        [
            ("A. Conserver React / Vinext", "Prévoir un hébergement compatible Cloudflare ou équivalent; réexaminer l'offre Hostinger."),
            ("B. Conserver Hostinger Premium", "Reproduire l'interface avec WordPress/PHP/MySQL et JavaScript pour la 3D et le motion."),
            ("C. Frontend statique + services", "Hostinger sert le site; formulaires, traduction, base et médias dépendent de services externes."),
        ],
        [2600, 6760],
    )
    add_body(doc, "À DÉCIDER : option d'architecture, responsabilité de maintenance, coûts des services externes et environnement de déploiement final.", highlight=True)

    missing_heading = add_h2(doc, "10.4 Informations manquantes")
    missing_heading.paragraph_format.page_break_before = True
    for item in (
        "Histoire officielle, mission, valeurs, chiffres de production et preuves commerciales.",
        "Pays déjà servis et pays recherchés pour de nouveaux représentants.",
        "Liste finale des produits, formats et caractéristiques publiables.",
        "Photographies des produits sous plusieurs angles et contenus de terrain.",
        "Coordonnées de contact, adresse, destinataires d'e-mails et informations juridiques.",
        "Liens officiels des réseaux sociaux.",
        "Règles de modération et budget de traduction.",
        "Architecture technique compatible avec l'hébergement retenu.",
    ):
        add_bullet(doc, item, highlight=True)

    # Appendix A: Prompt.
    add_h1(doc, "Annexe A. Prompt complet pour le professionnel UX/UI")
    add_callout(
        doc,
        "Mode d'utilisation",
        "Le texte ci-dessous peut être transmis tel quel à un consultant UX/UI ou à un agent de conception. Il complète le brief et exige une réponse structurée, critique et exploitable.",
        fill=GREEN_LIGHT,
        accent=GREEN,
    )
    add_h2(doc, "Prompt")
    prompt_paragraphs = [
        "Agis comme Lead Product Designer UX/UI et consultant motion design senior pour concevoir le futur site institutionnel et commercial de Multiproduit Mali SARL, domaine multiproduitmali.ml.",
        "Contexte : Multiproduit Mali est une entreprise malienne de commerce et de distribution proposant notamment les boissons fruitées Tropicoul et la boisson énergisante Triplex. Le nom SIMPARA était provisoire et ne doit apparaître nulle part. Le site doit d'abord convaincre des partenaires, importateurs, distributeurs et entrepreneurs susceptibles de représenter les produits dans d'autres pays. Le grand public et la communauté constituent une audience secondaire.",
        "Objectifs : établir rapidement la crédibilité de l'entreprise, valoriser le portefeuille, recruter des partenaires internationaux, publier les activités de MPM, permettre les réactions du public et rendre les contenus accessibles en français et en anglais.",
        "Direction souhaitée : équilibre 60 % institutionnel et 40 % spectaculaire. L'expérience doit être malienne, ambitieuse, généreuse, moderne et internationale. Le vert et l'orange MPM structurent l'identité; les pages produits utilisent les couleurs propres à chaque parfum. Le site ne doit pas ressembler à une landing page générique ni à un empilement de cartes.",
        "Hero recommandé : une parade fluide de produits. Le produit central devient actif et tourne lentement, tandis que les produits voisins restent partiellement visibles. Le scroll, le glissement tactile, les flèches et le clavier permettent de naviguer. La promesse et le bouton Devenir partenaire restent lisibles. Le hero laisse entrevoir la section suivante. Une carte vivante est utilisée plus bas pour les marchés et partenariats; les explosions de fruits, feuilles, glace et lumière sont réservées aux pages produits.",
        "Fonctions obligatoires : site FR/EN avec détection par préférence enregistrée, navigateur puis pays; choix manuel permanent; formulaire partenaire structuré avec courriels de notification; interface de gestion; actualités; commentaires modérés; traduction français-anglais des commentaires avec conservation de l'original; réseaux sociaux visibles; vrai responsive mobile; comportement reduced motion; solution de repli sans 3D.",
        "Ta mission : challenge les hypothèses du brief lorsque cela améliore la conversion ou la faisabilité, mais respecte le positionnement et les fonctions obligatoires. Ne te contente pas d'énumérer des tendances. Fournis des décisions précises, des mesures, des états d'interface et une justification orientée utilisateur et métier.",
        "Livrables attendus : 1) diagnostic des risques et opportunités; 2) architecture de l'information; 3) parcours partenaire, produit et communauté; 4) wireframes desktop et mobile; 5) deux variantes du hero avec recommandation; 6) direction visuelle et mini design system; 7) spécification motion détaillée; 8) maquette du formulaire et de ses états; 9) comportement de la langue et des traductions; 10) règles d'accessibilité et de performance; 11) handoff développeur avec tokens, breakpoints, tailles, espacements et logique d'animation; 12) liste des informations manquantes à demander au client.",
        "Contraintes de réussite : aucun chevauchement; pas de scroll-jacking; commandes tactiles d'au moins 44 x 44 px; navigation clavier; contraste WCAG 2.2 AA; texte lisible dans les deux langues; hero fonctionnel avant le chargement 3D; objectifs Core Web Vitals LCP <= 2,5 s, INP <= 200 ms et CLS <= 0,1; conception vérifiée à 360 x 800, 390 x 844, 768 x 1024, 1024 x 768 et 1440 x 900.",
        "Règles de présentation de ta réponse : commence par les décisions recommandées; illustre ensuite la page d'accueil section par section; sépare desktop et mobile; fournis une table de motion avec déclencheur, durée, easing, amplitude, repli reduced motion et coût en performance; termine par les questions bloquantes et les critères d'acceptation. Lorsque du code clarifie un comportement, donne un extrait court et directement exploitable, sans imposer une technologie avant validation de l'hébergement.",
    ]
    for idx, paragraph in enumerate(prompt_paragraphs):
        p = add_body(doc, paragraph, after=8)
        if idx == 0:
            p.runs[0].bold = True
            p.runs[0].font.color.rgb = rgb(GREEN_DARK)

    # Appendix B: Technical guidance.
    add_h1(doc, "Annexe B. Orientations techniques et extraits de code")
    add_body(doc, "Ces extraits sont des contrats d'implémentation indicatifs. Ils ne remplacent pas la validation de la pile d'hébergement. GSAP fonctionne avec différents frameworks; Three.js charge efficacement des modèles glTF/GLB et prend en charge Draco, Meshopt et KTX2 via ses chargeurs officiels.")

    css_code = r"""
:root {
  --mpm-green: #165b3a;
  --mpm-green-dark: #0f3e29;
  --mpm-orange: #f28a2b;
  --mpm-ink: #1c2722;
  --surface: #ffffff;
  --space: 8px;
  --motion-fast: 180ms;
  --motion-ui: 360ms;
  --motion-product: 800ms;
  --ease-out: cubic-bezier(.22, 1, .36, 1);
}

.hero {
  min-height: 86svh;
  display: grid;
  grid-template-columns: minmax(0, 5fr) minmax(18rem, 7fr);
  overflow: clip;
}

@media (max-width: 47.99rem) {
  .hero {
    grid-template-columns: 1fr;
    min-height: 88svh;
  }
}

@media (prefers-reduced-motion: reduce) {
  *, *::before, *::after {
    scroll-behavior: auto !important;
    animation-duration: 1ms !important;
    animation-iteration-count: 1 !important;
    transition-duration: 1ms !important;
  }
}
"""
    add_code_block(doc, css_code, "B.1 Tokens, hero responsive et réduction du mouvement")

    locale_code = r"""
type Locale = "fr" | "en";

const FRANCOPHONE = new Set([
  "ML", "SN", "CI", "BF", "NE", "TG", "BJ", "GN"
]);

export function resolveLocale(input: {
  saved?: Locale;
  browserLanguage?: string;
  countryCode?: string;
}): Locale {
  if (input.saved) return input.saved;
  const language = input.browserLanguage?.toLowerCase() ?? "";
  if (language.startsWith("fr")) return "fr";
  if (language.startsWith("en")) return "en";
  if (input.countryCode && FRANCOPHONE.has(input.countryCode)) return "fr";
  return "fr";
}
"""
    add_code_block(doc, locale_code, "B.2 Priorité de sélection de langue")

    translation_contract = r"""
type CommunityMessage = {
  id: string;
  originalText: string;
  originalLocale: "fr" | "en";
  translations: Partial<Record<"fr" | "en", {
    text: string;
    provider: string;
    generatedAt: string;
  }>>;
  moderationStatus: "pending" | "approved" | "hidden";
};

// L'interface affiche toujours l'original et charge la traduction à la demande.
// La traduction est enregistrée pour éviter un nouvel appel facturé.
"""
    add_code_block(doc, translation_contract, "B.3 Contrat de données pour un commentaire traduisible")

    gsap_code = r"""
import gsap from "gsap";
import { ScrollTrigger } from "gsap/ScrollTrigger";

gsap.registerPlugin(ScrollTrigger);

const mm = gsap.matchMedia();

mm.add("(min-width: 768px) and (prefers-reduced-motion: no-preference)", () => {
  const timeline = gsap.timeline({
    scrollTrigger: {
      trigger: "[data-product-parade]",
      start: "top 80%",
      end: "bottom 20%",
      scrub: 0.6
    }
  });

  timeline.to("[data-active-product]", {
    rotation: 360,
    xPercent: -10,
    ease: "none"
  });

  return () => timeline.scrollTrigger?.kill();
});
"""
    add_code_block(doc, gsap_code, "B.4 Exemple de liaison scroll-produit sans remplacer le scroll natif")

    gltf_code = r"""
import { GLTFLoader } from "three/addons/loaders/GLTFLoader.js";
import { DRACOLoader } from "three/addons/loaders/DRACOLoader.js";

const draco = new DRACOLoader();
draco.setDecoderPath("/draco/");

const loader = new GLTFLoader();
loader.setDRACOLoader(draco);

const model = await loader.loadAsync("/models/canette.glb");
scene.add(model.scene);
"""
    add_code_block(doc, gltf_code, "B.5 Chargement indicatif d'un GLB compressé")

    add_h2(doc, "B.6 Règles d'implémentation")
    implementation_rules = [
        "Ne charger Three.js et le modèle 3D que sur les pages qui l'utilisent.",
        "Ne pas masquer le H1, la promesse ou les actions derrière un canvas.",
        "Le canvas est décoratif pour les lecteurs d'écran; une image et un texte décrivent le produit.",
        "Prévoir une réserve de dimensions avant chargement pour éviter le déplacement de mise en page.",
        "Suspendre requestAnimationFrame lorsque le canvas est hors écran ou l'onglet inactif.",
        "Tester les interactions sur appareils tactiles et ne jamais imposer le drag comme seul mécanisme.",
    ]
    for rule in implementation_rules:
        p = add_bullet(doc, rule)
        p.paragraph_format.space_after = Pt(2)

    # Appendix C: Sources.
    add_h1(doc, "Annexe C. Références")
    add_body(doc, "Références consultées pour le cadrage. Elles servent de points de comparaison ou de documentation technique, et non de modèles à reproduire. Les décisions créatives finales restent à arbitrer après l'orientation UX/UI et la validation des éléments client encore manquants.")
    sources = [
        ("Ciao Energy - référence de tonalité produit", "https://www.ciaoenergy.org/"),
        ("SBC Kenya - portefeuille et parcours partenaire", "https://www.sbckenya.com/"),
        ("Equatorial Coca-Cola Bottling Company - présence et crédibilité institutionnelle", "https://www.eccbc.com/"),
        ("Meshy - bonnes pratiques Multi-View", "https://help.meshy.ai/en/articles/16102789-meshy-multi-view-best-practices-angles-and-images"),
        ("Meshy - produit et e-commerce", "https://help.meshy.ai/en/articles/16103073-meshy-for-product-design-and-e-commerce"),
        ("W3C - Web Content Accessibility Guidelines 2.2", "https://www.w3.org/TR/WCAG22/"),
        ("web.dev - seuils Core Web Vitals", "https://web.dev/articles/defining-core-web-vitals-thresholds"),
        ("GSAP - documentation ScrollTrigger", "https://gsap.com/docs/v3/Plugins/ScrollTrigger/"),
        ("Three.js - documentation GLTFLoader", "https://threejs.org/docs/pages/GLTFLoader.html"),
    ]
    for start in range(0, len(sources), 3):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.05
        for index, (label, url) in enumerate(sources[start:start + 3]):
            if index:
                separator = p.add_run("  |  ")
                set_font(separator, size=9.5, color=MUTED)
            add_hyperlink(p, label, url)

    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    path = build_document()
    print(path)
